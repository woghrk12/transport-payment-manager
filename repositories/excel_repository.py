import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from models.transport_record import TransportRecord


class ExcelRepository:
    SHEET_NAME_PATTERN = re.compile(r"^\d{4}\.(0[1-9]|1[0-2])")

    HEADERS = ("일", "회사명", "이름", "출발지", "목적지", "금액", "비고", "입금 날짜")

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)

    def get_all_records(self) -> list[TransportRecord]:
        if not self.file_path.is_file():
            raise FileNotFoundError(f"Excel 파일이 존재하지 않습니다: {self.file_path}")

        workbook = load_workbook(filename=self.file_path, read_only=True, data_only=True)

        records: list[TransportRecord] = []

        try:
            for worksheet in workbook.worksheets:
                if not self.SHEET_NAME_PATTERN.fullmatch(worksheet.title):
                    continue

                self._validate_headers(worksheet)

                for row_number, row in enumerate(worksheet.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
                    if (all(value is None or value == "" for value in row)):
                        continue
                    record = self._parse_record(worksheet.title, row_number, row)

                    records.append(record)
        finally:
            workbook.close()

        records.sort(key=lambda record: (record.work_date, record.sheet_name, record.row_number))

        return records

    def _validate_headers(self, worksheet) -> None:
        headers = tuple(worksheet.cell(row=1, column=column).value for column in range(1, 9))
        normalized = tuple(str(value).strip() if value is not None else "" for value in headers)

        if normalized != self.HEADERS:
            raise ValueError(f"'{worksheet.title} 시트의 열 구성이 예상한 형식과 다릅니다.")

    def _parse_record(self, sheet_name: str, row_number: int, row: tuple) -> TransportRecord:
        try:
            (work_date, company, name, departure, destination, amount, note, payment_date) = row

            if not company and not name:
                raise ValueError("회사와 이름이 모두 비어 있습니다.")

            return TransportRecord(
                sheet_name=sheet_name,
                row_number=row_number,
                work_date=self._parse_date(work_date),
                company=self._parse_text(company),
                name=self._parse_text(name),
                departure=self._parse_text(departure),
                destination=self._parse_text(destination),
                amount=self._parse_amount(amount),
                note=self._parse_text(note),
                payment_date=(self._parse_date(payment_date) if payment_date not in (None, "") else None)
            )
        except (ValueError, TypeError) as error:
            raise ValueError(f"{sheet_name} 시트 {row_number}행: {error}") from error

    @staticmethod
    def _parse_text(value) -> str:
        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def _parse_date(value) -> date:
        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, date):
            return value

        if isinstance(value, str):
            return date.fromisoformat(value.strip())

        raise ValueError(f"옳바르지 않은 날짜 형식입니다: {value}")

    @staticmethod
    def _parse_amount(value) -> int:
        if isinstance(value, bool):
            raise ValueError("금액은 숫자여야 합니다.")

        if isinstance(value, str):
            value = value.strip().replace(",", "")

        if value is None or value == "":
            raise ValueError("금액이 비어 있습니다.")

        try:
            amount = float(value)
        except (ValueError, TypeError) as error:
            raise ValueError(f"옳바르지 않은 금액입니다: {value}") from error

        if not amount.is_integer() or amount < 0:
            raise ValueError(f"금액은 0 이상의 정수여야 합니다: {value}")

        return int(amount)
