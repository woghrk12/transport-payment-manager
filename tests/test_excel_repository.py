from datetime import datetime

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from repositories.excel_repository import ExcelRepository


HEADERS = ["일", "회사명", "이름", "출발지", "목적지", "금액", "비고", "입금 날짜"]


@pytest.fixture
def sample_excel(tmp_path):
    file_path = tmp_path / "sample.xlsx"

    workbook = Workbook()

    september = workbook.active
    september.title = "2026.09"
    september.append(HEADERS)
    september.append([datetime(2026, 9, 20), "대한운수", "김철수", "서울", "부산", 300000, "", None])

    template: Worksheet = workbook.create_sheet("템플릿")
    template.append(["이 시트는 무시해야 합니다."])

    october: Worksheet = workbook.create_sheet("2026.10")
    october.append(HEADERS)
    october.append([datetime(2026, 10, 10), "대한운수", "김철수", "인천", "대전", 200000, "", None])
    october.append([datetime(2026, 10, 15), "한빛물류", "이영희", "대구", "광주", "250,000", "", datetime(2026, 11, 14)])

    workbook.save(file_path)
    workbook.close()
    return file_path


def test_reads_all_monthly_records(sample_excel):
    repository = ExcelRepository(str(sample_excel))

    records = repository.get_all_records()

    assert len(records) == 3
    assert records[0].sheet_name == "2026.09"
    assert records[0].row_number == 2
    assert records[1].sheet_name == "2026.10"
    assert records[2].row_number == 3


def test_converts_dates_and_amounts(sample_excel):
    repository = ExcelRepository(str(sample_excel))

    records = repository.get_all_records()

    assert records[0].work_date.isoformat() == "2026-09-20"
    assert records[0].amount == 300000
    assert records[0].payment_date is None

    assert records[2].amount == 250000
    assert records[2].payment_date.isoformat() == "2026-11-14"


def test_missing_file_raises_error(tmp_path):
    repository = ExcelRepository(str(tmp_path / "missing.xlsx"))

    with pytest.raises(FileNotFoundError):
        repository.get_all_records()


def test_invalid_headers_raise_error(tmp_path):
    file_path = tmp_path / "invalid.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "2026.10"

    worksheet.append(["날짜", "업체", "이름", "출발지", "목적지", "금액", "비고", "입금 날짜"])

    workbook.save(file_path)
    workbook.close()

    repository = ExcelRepository(str(file_path))

    with pytest.raises(ValueError, match="열 구성"):
        repository.get_all_records()
